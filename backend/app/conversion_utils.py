import os
import pdfkit
from docx2pdf import convert
from flask import jsonify
from pathlib import Path
from PIL import Image
from backend.app.file_utils import change_extension_to_pdf


def convert_txt_to_pdf(txt_file: str, output_pdf: str) -> None:
    """Convert a .txt file to PDF."""
    with open(txt_file, 'r') as file:
        text = file.read()
    pdfkit.from_string(f"<pre>{text}</pre>", output_pdf)

from docx import Document
import io
import mammoth
import pdfkit
from pathlib import Path


def convert_docx_to_pdf(docx_file: str, output_pdf: str) -> None:
    # Ensure paths are valid
    docx_path = Path(docx_file)
    output_path = Path(output_pdf)

    # Create a dummy Document if the file does not exist (you can remove this if you want to enforce existing files)
    if not docx_path.exists():
        document = Document()
        document.add_paragraph('Lorem ipsum dolor sit amet.')  # Placeholder content
        document.save(docx_path)

    # Read the .docx file
    document = Document(docx_path)
    
    # Create a bytestream for the document
    file_stream = io.BytesIO()
    document.save(file_stream)
    file_stream.seek(0)

    # Convert the docx to HTML
    result = mammoth.convert_to_html(file_stream)

    # Convert HTML to PDF
    pdf = pdfkit.from_string(result.value, False)  # False means we want to return the PDF as a byte string

    # Write the PDF to output file
    with open(output_path, 'wb') as file:
        file.write(pdf)


def convert_image_to_pdf(image_file: str, output_pdf: str) -> None:
    """Convert an image (jpeg, png) to PDF."""
    image = Image.open(image_file)
    image.save(output_pdf, "PDF")


def handle_pdf_conversion(uploaded_file_path: str) -> tuple:
    """Handles the PDF conversion logic for different file types."""
    try:
        # Get the file extension
        uploaded_file_path = Path(uploaded_file_path)
        file_extension = os.path.splitext(uploaded_file_path)[1].lower()
        pdf_output_path = Path(change_extension_to_pdf(uploaded_file_path))

        # Logic to handle different file types
        if file_extension == '.txt':
            convert_txt_to_pdf(uploaded_file_path, pdf_output_path)
        elif file_extension == '.html':
            pdfkit.from_file(uploaded_file_path, pdf_output_path)
        elif file_extension == '.docx':
            convert_docx_to_pdf(uploaded_file_path, pdf_output_path)
            breakpoint()
        elif file_extension in ['.jpeg', '.jpg', '.png']:
            convert_image_to_pdf(uploaded_file_path, pdf_output_path)
        else:
            return jsonify({"error": f"File type {file_extension} not supported."}), 400
        return jsonify({"message": "File converted successfully!", "file_path": str(pdf_output_path)}), 200

    except Exception as e:
        return jsonify({"error": f"PDF conversion failed: {str(e)}"}), 500
